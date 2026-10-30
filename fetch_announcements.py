"""讀取「116基金會補助整理表.xlsx」中的網址，抓取最新狀態，輸出 data.js 供網頁使用。

用法：
    python fetch_announcements.py [Excel 路徑]
"""
import datetime
import hashlib
import json
import re
import sys
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
DEFAULT_XLSX = Path(r"C:\Users\user\Downloads\116基金會補助整理表.xlsx")
STATE_FILE = HERE / "state.json"
DATA_FILE = HERE / "data.js"
HEADERS = {"User-Agent": "Mozilla/5.0 (fund-monitor; personal use)"}

URL_RE = re.compile(r"https?://[^\s]+")
DATE_RE = re.compile(r"(\d{2,4})年\s*(\d{1,2})月\s*(\d{1,2})日")
RANGE_END_RE = re.compile(r"\s*至\s*(\d{1,2})月\s*(\d{1,2})日")


def now_iso():
    return datetime.datetime.now().isoformat(timespec="minutes")


def clean(value):
    return " ".join(str(value or "").split())


def first_url(value):
    m = URL_RE.search(str(value or ""))
    return m.group(0).rstrip(")）,，") if m else ""


def parse_deadline(text):
    """取第一個完整日期；若為範圍（…至10月30日），取範圍結束日。民國年自動換算。"""
    m = DATE_RE.search(text or "")
    if not m:
        return None
    year = int(m.group(1))
    if year < 1911:
        year += 1911
    month, day = int(m.group(2)), int(m.group(3))
    end = RANGE_END_RE.match(text, m.end())
    if end:
        month, day = int(end.group(1)), int(end.group(2))
    try:
        return datetime.date(year, month, day).isoformat()
    except ValueError:
        return None


# 不在 Excel 表格中、但需要追蹤的基金會（沒有學校公告頁，因此以 note 顯示說明）
EXTRA_SOURCES = [
    {
        "name": "全聯佩樺圓夢社會福利基金會（大專校院服務性質社團支持補助計畫）",
        "deadline_text": "115年11月1日至11月15日（寒假階段，線上提案）",
        "phone": "02-8509-2688",
        "email": "Peihua@pxmart.com.tw",
        "need_doc": "線上提案（平台：青春有夢 投愛偏鄉）",
        "foundation_url": "https://www.phdf.org.tw/project/10/",
        "school_url": "",
        "note": "寒假階段每年 11/1–11/15 申請，寒假成果每年 3 月底前繳交；詳細規定見基金會官網。",
    },
    {
        "name": "金車文教基金會（2027金車曙光寒假服務物資贊助計劃）",
        "deadline_text": "115年12月15日 23:59 前（申請期間 11/1–12/15）",
        "phone": "(02)2100-1110",
        "email": "edu@kingcar.com.tw",
        "need_doc": "線上表單申請（需 Google 帳號登入）並上傳企劃書",
        "foundation_url": "https://kingcar.org.tw/project/501525",
        "school_url": "",
        "note": "服務時間 2027/1/5–2/28；審核結果於 2026/12/25 前公布；物資總值約 NT$3,000，寒暑假各 125 組。",
    },
    {
        "name": "國泰慈善基金會（國泰卓越獎助計畫）",
        "deadline_text": "請至官網確認",
        "phone": "",
        "email": "",
        "need_doc": "",
        "foundation_url": "https://patronc.cathaylife.com.tw/ODAY/F1/ODAYF100",
        "school_url": "",
        "note": "截止日與聯絡方式請至國泰慈善基金會官網確認。",
    },
]


def load_sources(xlsx_path):
    ws = load_workbook(xlsx_path, data_only=True).active
    sources = []
    for row in ws.iter_rows(min_row=3, max_col=7, values_only=True):
        name, deadline, phone, email, need_doc, found_url, school_url = row
        if not clean(name):
            continue
        deadline_text = clean(deadline)
        sources.append({
            "name": clean(name),
            "deadline_text": deadline_text,
            "deadline_iso": parse_deadline(deadline_text),
            "phone": clean(phone),
            "email": clean(email),
            "need_doc": clean(need_doc),
            "foundation_url": first_url(found_url),
            "school_url": first_url(school_url),
        })
    for extra in EXTRA_SOURCES:
        sources.append({**extra, "deadline_iso": parse_deadline(extra["deadline_text"])})
    return sources


def fetch(url):
    resp = requests.get(url, headers=HEADERS, timeout=30)
    resp.raise_for_status()
    resp.encoding = resp.apparent_encoding
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup


SPONSOR_KEYWORDS = ("贊助", "經費", "補助", "申請", "截止", "期限")


def sponsor_excerpt(content):
    """只保留內文中與社團贊助相關的句子。"""
    sentences = re.split(r"(?<=[。；])", content)
    picked = [s.strip() for s in sentences if any(k in s for k in SPONSOR_KEYWORDS)]
    return " ".join(picked)[:300]


def parse_school(soup):
    """學校公告頁（sso.nknu.edu.tw）：標題與社團贊助相關的內文摘錄。"""
    text = clean(soup.get_text(" "))

    def grab(pattern):
        m = re.search(pattern, text)
        return clean(m.group(1)) if m else ""

    title = grab(r"English\s+(.+?)\s+公告單位") or clean(soup.title.get_text() if soup.title else "")
    return {
        "title": title,
        "period": grab(r"公告時間[：:]\s*(.+?)\s+分類"),
        "excerpt": sponsor_excerpt(grab(r"點閱率[：:]\s*\d+\s+(.+?)\s+相關連結")),
    }


def digest(text):
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


def check_school(url, state, now):
    try:
        soup = fetch(url)
    except requests.RequestException as exc:
        return {"url": url, "ok": False, "error": str(exc)[:160]}

    info = parse_school(soup)
    fingerprint = digest(json.dumps(info, ensure_ascii=False, sort_keys=True))
    prev = state.get(url, {})
    changed_at = prev.get("changed_at")
    if prev.get("fingerprint") and prev["fingerprint"] != fingerprint:
        changed_at = now
    state[url] = {"fingerprint": fingerprint, "changed_at": changed_at}
    return {
        "url": url,
        "ok": True,
        **info,
        "changed_at": changed_at,
        "changed_now": changed_at == now,
    }


def check_foundation(url):
    try:
        soup = fetch(url)
    except requests.RequestException as exc:
        return {"url": url, "ok": False, "error": str(exc)[:160]}
    return {"url": url, "ok": True, "title": clean(soup.title.get_text() if soup.title else "")}


# 高師大「公告系統」分類列表（含發布單位欄）；以發布單位與關鍵字篩出課外活動指導組的社團贊助公告
NKNU_LIST_URL = "https://sso.nknu.edu.tw/News/SubjectList.aspx?cat=CAMPUS"
HIDDEN_RE = re.compile(r'<input[^>]*type="hidden"[^>]*name="([^"]+)"[^>]*value="([^"]*)"')
NOTICE_UNIT = "課外活動指導組"
NOTICE_KEYWORDS = ("社團", "贊助", "補助", "經費", "服務隊")
NOTICE_MAX_PAGES = 30
NOTICE_HISTORY_FILE = HERE / "notices_history.json"


def parse_notice_rows(html):
    """逐列解析（每列以分類欄位分隔），避免標題含標籤時跨列誤配。"""
    rows = []
    for chunk in html.split('_Category">')[1:]:
        cat = re.search(r'^([^<]*)</span>', chunk)
        link = re.search(r'linkSubject_short"[^>]*?href="([^"]+)"[^>]*>(.*?)</a>', chunk, re.S)
        dept = re.search(r'_DeptName">([^<]*)</span>', chunk)
        start = re.search(r'_StartDate">([^<]*)</span>', chunk)
        end = re.search(r'_EndDate">([^<]*)</span>', chunk)
        if not (cat and link and dept and start and end):
            continue
        rows.append({
            "category": cat.group(1).strip(),
            "url": link.group(1).strip(),
            "title": clean(re.sub(r"<[^>]+>", "", link.group(2))),
            "unit": clean(dept.group(1)),
            "start": start.group(1).strip(),
            "end": end.group(1).strip(),
        })
    return rows


def crawl_notice_rows():
    """依序翻頁抓取公告列表（ASP.NET GridView 分頁為 postback）。"""
    session = requests.Session()
    session.headers.update(HEADERS)
    resp = session.get(NKNU_LIST_URL, timeout=30)
    resp.raise_for_status()
    resp.encoding = resp.apparent_encoding

    rows = []
    for page in range(1, NOTICE_MAX_PAGES + 1):
        rows.extend(parse_notice_rows(resp.text))
        if f"Page${page + 1}" not in resp.text:
            break
        body = {name: value for name, value in HIDDEN_RE.findall(resp.text)}
        body.update({"__EVENTTARGET": "gvFull", "__EVENTARGUMENT": f"Page${page + 1}"})
        resp = session.post(NKNU_LIST_URL, data=body, timeout=30)
        resp.raise_for_status()
        resp.encoding = resp.apparent_encoding
    return rows


def load_notice_history():
    """讀取已保存的歷史公告（以 UNID 為鍵）；檔案不存在代表第一次執行。"""
    if NOTICE_HISTORY_FILE.exists():
        return json.loads(NOTICE_HISTORY_FILE.read_text(encoding="utf-8")), False
    return {}, True


def fetch_club_notices(state, now):
    """回傳課外活動指導組的社團贊助公告（歷史累積），並標記本次新增者。"""
    history, first_run = load_notice_history()
    warning = ""
    try:
        rows = crawl_notice_rows()
    except requests.RequestException as exc:
        print(f"公告列表讀取失敗，改用歷史公告：{exc}")
        warning = str(exc)[:160]
        rows = []

    new_unids = set()
    for row in rows:
        if NOTICE_UNIT not in row["unit"] or not any(k in row["title"] for k in NOTICE_KEYWORDS):
            continue
        unid = row["url"].split("UNID=")[-1]
        if unid not in history:
            history[unid] = {**row, "unid": unid, "first_seen": now}
            if not first_run:
                new_unids.add(unid)
        else:
            # 列表內容可能被學校修改，以最新資料覆蓋，但保留首次發現時間
            history[unid] = {**row, "unid": unid, "first_seen": history[unid]["first_seen"]}

    NOTICE_HISTORY_FILE.write_text(json.dumps(history, ensure_ascii=False, indent=2), encoding="utf-8")

    items = sorted(history.values(), key=lambda item: item["start"], reverse=True)
    for item in items:
        item["is_new"] = item["unid"] in new_unids
    print(f"社團贊助公告：歷史共 {len(items)} 筆，本次新增 {len(new_unids)} 筆")
    return {"ok": not warning or bool(items), "checked": now, "items": items, "warning": warning}


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    now = now_iso()
    sources = load_sources(xlsx)
    state = json.loads(STATE_FILE.read_text(encoding="utf-8")) if STATE_FILE.exists() else {}

    for src in sources:
        src["school"] = check_school(src["school_url"], state, now) if src["school_url"] else None
        src["foundation"] = check_foundation(src["foundation_url"]) if src["foundation_url"] else None
        print(f"{src['name'][:20]:<20} 截止 {src['deadline_iso']} | "
              f"學校頁 {'-' if not src['school'] else ('OK' if src['school']['ok'] else 'ERR')} | "
              f"基金會 {'-' if not src['foundation'] else ('OK' if src['foundation']['ok'] else 'ERR')}")

    notices = fetch_club_notices(state, now)

    payload = json.dumps({"generated_at": now, "sources": sources, "club_notices": notices}, ensure_ascii=False, indent=2)
    # 避免內容中的 </script> 提前結束標籤
    payload = payload.replace("</", "<\\/")
    DATA_FILE.write_text("window.FUND_DATA = " + payload + ";\n", encoding="utf-8")
    STATE_FILE.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"已輸出 {DATA_FILE}")


if __name__ == "__main__":
    main()
